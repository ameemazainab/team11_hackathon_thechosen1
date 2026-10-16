"""FR-14, FR-15, FR-16: snapshots, exports, actuals and the accuracy report; insights."""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app import models as M
from app.domain.util import m_label
from app.jobs import monthly_snapshot
from app.models import ImmutableRecordError


def sheet_rows(content: bytes):
    ws = load_workbook(io.BytesIO(content)).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def find(rows, label):
    return next(r for r in rows if label in r)


def test_fr14_snapshot_immutable_after_later_edits(as_user, db):
    fin, ac = as_user("FIN"), as_user("AC")
    r = fin.post("/api/snapshots", json={"label": "Before the call"})
    assert r.status_code == 201
    sid = r.json()["id"]
    before = fin.get(f"/api/snapshots/{sid}").json()
    assert before["committed"]["2026-11"] == 2_813_000 and before["weighted"]["2026-11"] == 164_713
    assert ac.patch("/api/projects/R1001/months/2026-11", json={"value": 99000}).status_code == 200
    assert fin.get("/api/forecast").json()["totals"]["committed"]["2026-11"] == 2_832_000  # live data moved
    after = fin.get(f"/api/snapshots/{sid}").json()
    assert after == before  # the snapshot did not
    # and the database refuses to change or delete one
    snap = db.query(M.Snapshot).filter_by(id=sid).one()
    snap.label = "tampered"
    with pytest.raises(ImmutableRecordError):
        db.flush()
    db.rollback()
    with pytest.raises(ImmutableRecordError):
        db.delete(db.query(M.Snapshot).filter_by(id=sid).one())
        db.flush()
    db.rollback()
    assert fin.get("/api/snapshots").json()[0]["id"] == sid


def test_snapshot_permissions(as_user):
    assert as_user("AC").post("/api/snapshots", json={}).status_code == 403
    assert as_user("CRO").post("/api/snapshots", json={}).status_code == 403
    assert as_user("ADM").post("/api/snapshots", json={}).status_code == 201


def test_fr15_export_totals_equal_on_screen_xlsx_and_csv(as_user):
    cro = as_user("CRO")
    f = cro.get("/api/forecast", params={"region": "NA", "tech": "PEGA"}).json()
    r = cro.get("/api/export/forecast", params={"format": "xlsx", "region": "NA", "tech": "PEGA"})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    rows = sheet_rows(r.content)
    flat = [str(c) for row in rows[:8] for c in row if c]
    assert any("Exported" == c for c in flat) and any(c.startswith("2026-10-01") for c in flat)  # export timestamp
    assert any("Region: NA; Practice: PEGA" in c or "Practice: PEGA" in c for c in flat)  # filters applied
    header = next(row for row in rows if row and row[0] == "Row ID")
    for label, key in (("Committed total", "committed"), ("Weighted pipeline total", "weighted"), ("Total forecast", "total")):
        tot = find(rows, label)
        for i, h in enumerate(header):
            if h == "Nov 26":
                assert tot[i] == round(f["totals"][key]["2026-11"]), label
            if h == "Mar 27":
                assert tot[i] == round(f["totals"][key]["2027-03"]), label
    csv = cro.get("/api/export/forecast", params={"format": "csv", "region": "NA", "tech": "PEGA"})
    assert csv.status_code == 200 and "Filters applied" in csv.text and "Total forecast" in csv.text
    assert cro.get("/api/export/forecast", params={"format": "pdf"}).status_code == 422


def test_export_rows_match_grid_rows(as_user):
    cro = as_user("CRO")
    f = cro.get("/api/forecast").json()
    rows = sheet_rows(cro.get("/api/export/forecast").content)
    header = next(r for r in rows if r and r[0] == "Row ID")
    ids = [r[0] for r in rows if r and isinstance(r[0], str) and r[0].startswith("R") and r[0][1:].isdigit()]
    assert sorted(ids) == sorted(x["id"] for x in f["rows"])
    i = header.index("Oct 26")
    assert sum(r[i] or 0 for r in rows if r and r[0] in set(ids)) == f["totals"]["committed"]["2026-10"]


def test_snapshot_export_totals_equal_snapshot_view(as_user):
    fin = as_user("FIN")
    sid = fin.post("/api/snapshots", json={"label": "Export test"}).json()["id"]
    as_user("AC").patch("/api/projects/R1001/months/2026-12", json={"value": 5})  # later edit must not leak into the export
    view = fin.get(f"/api/snapshots/{sid}").json()
    rows = sheet_rows(fin.get(f"/api/snapshots/{sid}/export", params={"format": "xlsx"}).content)
    header = next(r for r in rows if r and r[0] == "Row ID")
    tot = find(rows, "Total committed")
    for i, h in enumerate(header):
        if h and h[:3] in ("Oct", "Dec", "Mar"):
            m = next(mm for mm in view["months"] if m_label(mm) == h)
            assert tot[i] == view["committed"][m], h
    assert any("Export test" in str(c) for r in rows[:6] for c in r)


def test_seed_snapshot_listed_and_compare(as_user):
    fin = as_user("FIN")
    snaps = fin.get("/api/snapshots").json()
    assert snaps[-1]["id"] == "snap-2026-09-01" and snaps[-1]["auto"] is True and snaps[-1]["rows"] == 109
    mv = fin.get("/api/insights/movement").json()
    top = mv["movers"][0]
    assert top["row_id"] == "R1112" or top["customer"]
    kestrel = next(m for m in mv["movers"] if m["row_id"] == "R1001")
    assert kestrel["delta"] < 0 and kestrel["reason"] == "SOW ending / ramp-down"
    harbour = next(m for m in mv["movers"] if m["row_id"] == "R1002")
    assert harbour["reason"] is None and mv["unexplained"] >= 1
    assert mv["total_now"] != mv["total_was"]


def test_monthly_snapshot_job_first_business_day(db, env):
    # 2026-10-01 is a Thursday: the first business day of October
    assert monthly_snapshot.first_business_day(date(2026, 10, 1)).day == 1
    assert monthly_snapshot.first_business_day(date(2026, 8, 15)).day == 3  # Aug 1 is a Saturday
    r = monthly_snapshot.job(db, env)
    assert r["ok"] and r["snapshot"].startswith("snap-2026-10-01")
    assert "already exists" in monthly_snapshot.job(db, env)["skipped"]
    env.as_of_date = "2026-10-15"
    assert "not the first business day" in monthly_snapshot.job(db, env)["skipped"]
    env.as_of_date = "2026-10-01"
    snap = db.query(M.Snapshot).filter(M.Snapshot.date == "2026-10-01").one()
    assert snap.auto and snap.taken_by == "Scheduled job" and snap.label == "Month-start snapshot, October 2026"


CSV = "Customer,Project,Amount,Month\nKestrel Life & Annuity,Policy Admin Modernization,80000,2026-12\nNorthwind Mutual,Claims Platform,2400000,Dec 2026\nBad Co,X,abc,2026-12\n,Y,5,2026-12\n"


def test_fr16_actuals_upload_and_accuracy_vs_m3_snapshot(as_user):
    fin = as_user("FIN")
    assert as_user("AC").post("/api/actuals/upload", files={"file": ("a.csv", CSV, "text/csv")}).status_code == 403
    r = fin.post("/api/actuals/upload", files={"file": ("a.csv", CSV, "text/csv")})
    assert r.status_code == 200
    body = r.json()
    assert body["loaded"] == 2 and len(body["rejects"]) == 2 and body["months"] == ["2026-12"]
    assert {x["line"] for x in body["rejects"]} == {4, 5}
    actual = 2_480_000
    rep = fin.get("/api/insights/accuracy").json()
    dec = next(m for m in rep["months"] if m["month"] == "2026-12")
    assert dec["actual"] == actual
    m3 = next(k for k in dec["lookaheads"] if k["k"] == 3)
    assert m3["snapshot_id"] == "snap-2026-09-01"
    assert m3["committed"] == 2_432_000  # Dec on the 1 Sep snapshot
    assert m3["committed_variance"] == 2_432_000 - actual and m3["committed_variance_pct"] == round((2_432_000 - actual) / actual * 100, 2)
    assert m3["committed_flag"] is False  # -1.94% is inside +/-5%
    assert next(k for k in dec["lookaheads"] if k["k"] == 2)["snapshot_id"] is None  # no Oct snapshot yet
    # a smaller actual puts it outside the tolerance and flags it
    csv2 = "Customer,Project,Amount,Month\nNorthwind Mutual,Claims Platform,2000000,2026-12\n"
    fin.post("/api/actuals/upload", files={"file": ("b.csv", csv2, "text/csv")})
    dec2 = next(m for m in fin.get("/api/insights/accuracy").json()["months"] if m["month"] == "2026-12")
    k3 = next(k for k in dec2["lookaheads"] if k["k"] == 3)
    assert dec2["actual"] == 2_000_000 and k3["committed_flag"] is True and k3["committed_variance_pct"] == 21.6


def test_actuals_xlsx_and_bad_files(as_user):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Customer", "Project", "Amount", "Month"])
    ws.append(["Kestrel Life & Annuity", "Policy Admin Modernization", 80000, "2026-11"])
    buf = io.BytesIO()
    wb.save(buf)
    fin = as_user("FIN")
    r = fin.post("/api/actuals/upload", files={"file": ("a.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 200 and r.json()["loaded"] == 1 and r.json()["unmatched"] == []
    assert fin.post("/api/actuals/upload", files={"file": ("a.csv", "Foo,Bar\n1,2\n", "text/csv")}).status_code == 422
    assert fin.post("/api/actuals/upload", files={"file": ("a.pdf", "x", "application/pdf")}).status_code == 422
    assert fin.get("/api/actuals").json()[0]["month"] == "2026-11"


def test_insights_views(as_user):
    res = as_user("RES")  # Resourcing: read-only demand view by practice and month (FR-22)
    d = res.get("/api/insights/by", params={"dim": "tech"}).json()
    assert d["rows"][0]["key"] and d["months"][0] == "2026-10" and d["rows"][0]["weighted"] is not None
    assert sum(r["committed"]["2026-10"] for r in d["rows"]) == 2_677_000
    assert res.get("/api/insights/by", params={"dim": "nope"}).status_code == 422
    gap = res.get("/api/insights/fy-gap").json()
    assert gap["available"] and gap["fy_end"] == "2027-03" and gap["baseline"] == 2_780_000
    assert gap["series"][1]["month"] == "2027-04" and gap["series"][1]["committed"] == 2_192_250
    assert gap["stopping_total"] > 0 and len(gap["stopping"]) >= 6
